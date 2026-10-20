"""Excel syllabus parser.

Turns a heterogeneous syllabus sheet into classified rows + a course tree, driven by a
*mapping profile* (column mapping + classification rules). Nothing is hardcoded per file:
the three presets below are just profiles, and teachers can edit any of them.

Row types: TITLE, HEADER, GROUP_L1 (Học phần / Phần), GROUP_L2 (Module),
SESSION, PROJECT, NOTE, IGNORE.
"""
from __future__ import annotations

import io
import re
from copy import deepcopy
from dataclasses import asdict, dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

DEFAULT_RULES = {
    "group_l1": r"^\s*(Học phần|HỌC PHẦN|PHẦN|Phần)\s*\d+",
    "group_l2": r"^\s*Module\s*\d+",
    "session": r"^\W*Buổi\s*[\d]",
    "project": r"^\s*Project\b",
    "note_min_chars": 120,
    "practice_keywords": ["thực hành", "bài tập", "project", "code:", "yêu cầu:"],
}

PRESETS: dict[str, dict] = {
    "ai": {
        "label": "AI Engineer (Module → Buổi; mục tiêu gộp dọc theo module)",
        "header_row": 3,
        "columns": {"label": "A", "title": "B", "content": "C", "objectives": "D"},
        "title_from_label": False,
        "objectives_prefix": None,
        "rules": DEFAULT_RULES,
    },
    "de": {
        "label": "Data Engineer (Học phần → Module → Buổi/Project)",
        "header_row": 3,
        "columns": {"group": "A", "label": "B", "title": "C", "content": "D"},
        "title_from_label": False,
        "objectives_prefix": None,
        "rules": DEFAULT_RULES,
    },
    "ba": {
        "label": "Business Analyst (Phần → Module → Buổi; thực hành cột riêng)",
        "header_row": 3,
        "columns": {"label": "A", "objectives": "B", "content": "C", "practice": "D"},
        "title_from_label": True,
        "objectives_prefix": "Mục tiêu:",
        "rules": DEFAULT_RULES,
    },
}


@dataclass
class ParsedRow:
    row: int
    type: str
    label: str = ""
    title: str = ""
    content: str = ""
    objectives: str = ""
    practice: str = ""
    session_numbers: list[int] = field(default_factory=list)
    objectives_scope: str = "lesson"  # lesson | module (vertically merged cell)
    warnings: list[str] = field(default_factory=list)


def _clean(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\r\n", "\n").replace("\r", "\n")
    lines = [ln.rstrip() for ln in s.split("\n")]
    return "\n".join(lines).strip()


def _first_line(s: str) -> str:
    return s.split("\n", 1)[0].strip()


_SESSION_NUM = re.compile(r"Buổi\s*([\d][\d\s,+]*)", re.IGNORECASE)


def parse_session_label(label: str) -> tuple[list[int], str]:
    """'Buổi 42,43' -> ([42,43], ''); '★ Buổi 21 AI Mastery: X' -> ([21], 'AI Mastery: X');
    'Buổi 1: Nghề BA' -> ([1], 'Nghề BA'); 'Buổi 1+2' -> ([1,2], '')."""
    first = _first_line(label)
    m = _SESSION_NUM.search(first)
    if not m:
        return [], first
    nums = [int(x) for x in re.findall(r"\d+", m.group(1))]
    rest = first[m.end():].strip().lstrip(":–-").strip()
    return nums, rest


def _merged_lookup(ws):
    """Map (row, col) -> merged range for every cell covered by a merged range."""
    cover = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                cover[(r, c)] = rng
    return cover


def list_sheets(data: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def detect_preset(data: bytes) -> str | None:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.worksheets[0]
    header = " | ".join(_clean(c.value) for c in ws[3])
    title = _clean(ws["A1"].value).upper()
    wb.close()
    if "Học phần" in header:
        return "de"
    if "Thực hành" in header and "Chủ đề" in header:
        return "ba"
    if "AI ENGINEER" in title or "Mục tiêu" in header:
        return "ai"
    return None


def parse_workbook(data: bytes, profile: dict, sheet: str | None = None) -> dict:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    cover = _merged_lookup(ws)
    rules = {**DEFAULT_RULES, **(profile.get("rules") or {})}
    re_l1 = re.compile(rules["group_l1"])
    re_l2 = re.compile(rules["group_l2"])
    re_sess = re.compile(rules["session"], re.IGNORECASE)
    re_proj = re.compile(rules["project"], re.IGNORECASE)
    cols = {k: column_index_from_string(v) for k, v in profile["columns"].items() if v}
    header_row = int(profile.get("header_row", 1))
    title_from_label = bool(profile.get("title_from_label"))
    obj_prefix = profile.get("objectives_prefix")
    practice_kw = [k.lower() for k in rules["practice_keywords"]]

    def cell(r: int, key: str) -> str:
        c = cols.get(key)
        if not c:
            return ""
        rng = cover.get((r, c))
        if rng is not None and (rng.min_row, rng.min_col) != (r, c):
            return ""  # covered by a merge whose anchor is elsewhere
        return _clean(ws.cell(r, c).value)

    rows: list[ParsedRow] = []
    title_text = _clean(ws.cell(1, 1).value)
    for r in range(1, ws.max_row + 1):
        values = [_clean(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        non_empty = [v for v in values if v]
        if not non_empty:
            continue
        if r < header_row:
            rows.append(ParsedRow(row=r, type="TITLE" if r == 1 else "IGNORE", title=non_empty[0]))
            continue
        if r == header_row:
            rows.append(ParsedRow(row=r, type="HEADER", title=" / ".join(non_empty)))
            continue

        label = cell(r, "label") or cell(r, "group")
        group_cell = cell(r, "group")
        lead = group_cell or label or non_empty[0]
        pr = ParsedRow(row=r, type="IGNORE", label=_first_line(label))

        if group_cell and re_l1.search(group_cell):
            pr.type, pr.title = "GROUP_L1", _first_line(group_cell)
        elif re_l1.search(lead):
            pr.type, pr.title = "GROUP_L1", _first_line(lead)
        elif re_l2.search(lead) or (label and re_l2.search(label)):
            pr.type, pr.title = "GROUP_L2", _first_line(label or lead)
        elif label and (re_sess.search(label) or re_proj.search(label)):
            pr.type = "PROJECT" if re_proj.search(label) else "SESSION"
            nums, rest = parse_session_label(label)
            pr.session_numbers = nums
            title = rest if title_from_label else cell(r, "title")
            content = cell(r, "content")
            objectives = cell(r, "objectives")
            practice = cell(r, "practice")
            if obj_prefix and objectives.startswith(obj_prefix):
                objectives = objectives[len(obj_prefix):].strip()
            # Vertically merged objectives cell => belongs to the module, not the lesson.
            oc = cols.get("objectives")
            if oc and (rng := cover.get((r, oc))) is not None and rng.max_row > rng.min_row:
                pr.objectives_scope = "module"
            if not title:
                title = _first_line(objectives) or _first_line(content) or rest or pr.label
                if title == _first_line(objectives):
                    pr.warnings.append("TITLE_FROM_OBJECTIVES")
            pr.title, pr.content, pr.objectives, pr.practice = title, content, objectives, practice
            if len(nums) > 1:
                pr.warnings.append("MULTI_SESSION_ROW")
            if re.fullmatch(r"[\W\w]{0,3}", title) or not re.search(r"[A-Za-zÀ-ỹ]{3,}", title):
                pr.warnings.append("SUSPICIOUS_TITLE")
        elif len(non_empty) == 1 and len(non_empty[0]) >= int(rules["note_min_chars"]):
            pr.type, pr.title = "NOTE", non_empty[0]
        elif len(non_empty) == 1:
            pr.type, pr.title = "GROUP_L1", _first_line(non_empty[0])
            pr.warnings.append("GROUP_INFERRED")
        else:
            pr.title = _first_line(non_empty[0])
            pr.warnings.append("UNCLASSIFIED")
        rows.append(pr)
    wb.close()

    result = {
        "sheet": ws.title,
        "title": title_text,
        "rows": [asdict(x) for x in rows],
    }
    result.update(build_tree(result["rows"], title_text))
    return result


def build_tree(rows: list[dict], course_title: str = "") -> dict:
    """Assemble modules/lessons from (possibly teacher-edited) classified rows."""
    modules: list[dict] = []
    warnings: list[dict] = []
    current_l1: str | None = None
    current_mod: dict | None = None
    seen: dict[tuple, int] = {}

    def new_module(title: str, implicit: bool, row: int) -> dict:
        m = re.match(r"^\s*Module\s*(\d+)", title)
        code = f"M{m.group(1)}" if m else f"G{len(modules) + 1}"
        mod = {"code": code, "title": title, "group_path": current_l1, "objectives": "",
               "lessons": [], "implicit": implicit, "row": row}
        modules.append(mod)
        return mod

    for r in rows:
        t = r["type"]
        for w in r.get("warnings") or []:
            warnings.append({"row": r["row"], "code": w, "text": r.get("title", "")[:80]})
        if t == "GROUP_L1":
            current_l1, current_mod = r["title"], None
        elif t == "GROUP_L2":
            current_mod = new_module(r["title"], False, r["row"])
        elif t in ("SESSION", "PROJECT"):
            if current_mod is None:
                current_mod = new_module(current_l1 or "Chưa phân nhóm", True, r["row"])
                warnings.append({"row": r["row"], "code": "IMPLICIT_MODULE", "text": current_mod["title"][:80]})
            key = (current_l1, tuple(r["session_numbers"]), r["title"].strip().lower())
            if key in seen:
                warnings.append({"row": r["row"], "code": "DUPLICATE_ROW",
                                 "text": f"Trùng với hàng {seen[key]}: {r['title'][:60]}"})
                continue
            seen[key] = r["row"]
            lesson_objectives = r.get("objectives") or ""
            if r.get("objectives_scope") == "module" and lesson_objectives:
                if not current_mod["objectives"]:
                    current_mod["objectives"] = lesson_objectives
                lesson_objectives = ""
            text_blob = " ".join([r.get("content") or "", r.get("practice") or "", r["title"]]).lower()
            current_mod["lessons"].append({
                "label": r.get("label") or "",
                "session_numbers": r["session_numbers"],
                "title": r["title"],
                "content": r.get("content") or "",
                "objectives": lesson_objectives,
                "practice": r.get("practice") or "",
                "kind": "project" if t == "PROJECT" else "session",
                "practice_candidate": t == "PROJECT" or bool(r.get("practice"))
                or any(k in text_blob for k in DEFAULT_RULES["practice_keywords"]),
                "row": r["row"],
            })
    session_count = len({(m["group_path"] if m["implicit"] else None, n)
                         for m in modules for ls in m["lessons"] for n in ls["session_numbers"]})
    return {
        "course_title": course_title,
        "modules": modules,
        "warnings": warnings,
        "stats": {
            "modules": sum(1 for m in modules if not m["implicit"]),
            "groups": len({m["group_path"] for m in modules if m["group_path"]}),
            "lessons": sum(len(m["lessons"]) for m in modules),
            "distinct_session_numbers": session_count,
        },
    }


def preset(name: str) -> dict:
    return deepcopy(PRESETS[name])
